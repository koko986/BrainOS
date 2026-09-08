from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "presentation" / "MARLIN_Easy_Project_Guide.docx"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def borders(table, colour: str = "D9D9D9") -> None:
    properties = table._tbl.tblPr
    element = properties.find(qn("w:tblBorders"))
    if element is None:
        element = OxmlElement("w:tblBorders")
        properties.append(element)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "6")
        item.set(qn("w:color"), colour)
        element.append(item)


def cell_margins(cell, top: int = 110, start: int = 130, bottom: int = 110, end: int = 130) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, name: str = "Aptos", size: float | None = None, bold: bool | None = None, colour: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if colour:
        run.font.color.rgb = RGBColor.from_string(colour)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, *, bold_lead: str = "") -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead:
        set_font(paragraph.add_run(bold_lead), bold=True)
    set_font(paragraph.add_run(text))


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        set_font(paragraph.add_run(item))


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.28)
        paragraph.paragraph_format.space_after = Pt(4)
        set_font(paragraph.add_run(f"{index}.  {item}"))


def add_flow(doc: Document, steps: list[str]) -> None:
    table = doc.add_table(rows=1, cols=len(steps) * 2 - 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    usable = 6.65
    arrow_width = 0.28
    box_width = (usable - arrow_width * (len(steps) - 1)) / len(steps)
    for index, cell in enumerate(table.rows[0].cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(cell, 95, 70, 95, 70)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if index % 2:
            cell.width = Inches(arrow_width)
            set_font(paragraph.add_run("->"), size=11, bold=True, colour="366B72")
        else:
            cell.width = Inches(box_width)
            shade(cell, "E9F3F2")
            set_font(paragraph.add_run(steps[index // 2]), size=9.5, bold=True, colour="16343A")
    borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_component_table(doc: Document) -> None:
    rows = [
        ("Qwen3 4B", "Conversation brain", "Understands questions and writes short replies"),
        ("Python", "Controller", "Routes commands and performs Windows actions"),
        ("SQLite", "Memory", "Stores knowledge, history, reminders, and relationships"),
        ("Prolog", "Logical thinker", "Applies rules and explains task recommendations"),
        ("Brain graph", "Knowledge view", "Shows projects, folders, files, and connections"),
        ("Whisper and Vosk", "Ears", "Transcribes speech and detects Hey MARLIN"),
        ("Piper", "Voice", "Speaks English replies locally"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (1.35, 1.35, 3.95)
    headers = ("Technology", "Simple role", "Function")
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        shade(cell, "173B57")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(cell)
        set_font(cell.paragraphs[0].add_run(text), size=10, bold=True, colour="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, text in enumerate(values):
            cells[column].width = Inches(widths[column])
            cells[column].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cells[column])
            if row_index % 2:
                shade(cells[column], "F3F7FA")
            set_font(cells[column].paragraphs[0].add_run(text), size=9.5)
    borders(table)


def add_code(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "F2F4F5")
    cell_margins(cell, 130, 180, 130, 180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line + ("\n" if index < len(lines) - 1 else ""))
        set_font(run, "Consolas", 9)
    borders(table, "D1D6D8")


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    for name in ("Normal", "List Bullet", "List Number"):
        styles[name].font.name = "Aptos"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        styles[name].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(30)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_style_properties = styles["Title"]._element.get_or_add_pPr()
    title_style_border = title_style_properties.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_properties.remove(title_style_border)
    for level, size in ((1, 18), (2, 13)):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(30)
    title.paragraph_format.space_after = Pt(8)
    title.add_run("MARLIN Local AI Assistant")
    title_properties = title._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(22)
    set_font(subtitle.add_run("Easy Project Guide and Presentation Notes"), size=16, colour="366B72")
    add_body(doc, "MARLIN is a private personal AI assistant that runs on a Windows laptop. It combines local conversation, voice control, computer actions, reminders, a visual knowledge graph, and Prolog logical reasoning without requiring a paid API key.")
    add_body(doc, "Main conclusion: ", bold_lead="")
    conclusion = doc.paragraphs[-1]
    conclusion.clear()
    set_font(conclusion.add_run("Main conclusion: "), bold=True)
    set_font(conclusion.add_run("MARLIN is more than a chatbot because it can remember information, explain logical decisions, and perform real actions on the laptop."))

    add_heading(doc, "The Main Idea")
    add_body(doc, "The easiest way to understand MARLIN is to compare its software components to parts of a person.")
    add_component_table(doc)

    doc.add_page_break()
    add_heading(doc, "System Architecture")
    add_flow(doc, ["Speak or type", "Desktop cockpit", "Python router", "Choose service", "Return result"])
    add_body(doc, "Python is the central controller. It decides whether a request should use a direct Windows action, Prolog reasoning, or the Qwen language model.")
    add_heading(doc, "Three Main Routes", 2)
    add_bullets(doc, [
        "Computer route: open camera, open an application, play a video, or search a file through predefined Python actions.",
        "Reasoning route: load task facts from SQLite, apply Prolog rules, and return a conclusion with an explanation.",
        "Conversation route: send ordinary conversation to the local Qwen3 4B model and stream the reply to the cockpit and voice system.",
    ])
    add_heading(doc, "Command Examples", 2)
    examples = [
        ("open camera", "Python opens the Windows Camera app immediately."),
        ("play C:\\Videos\\demo.mp4", "Python opens the local file in the default video player."),
        ("high priority tasks", "Prolog applies priority rules to task facts."),
        ("hello MARLIN", "Qwen produces a conversational reply."),
        ("delete this file", "MARLIN shows the exact target and asks for confirmation."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, text in enumerate(("Command", "What happens")):
        shade(table.rows[0].cells[index], "173B57")
        cell_margins(table.rows[0].cells[index])
        set_font(table.rows[0].cells[index].paragraphs[0].add_run(text), size=10, bold=True, colour="FFFFFF")
    for row_index, (command, result) in enumerate(examples):
        cells = table.add_row().cells
        if row_index % 2:
            shade(cells[0], "F3F7FA")
            shade(cells[1], "F3F7FA")
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
        set_font(cells[0].paragraphs[0].add_run(command), "Consolas", 9)
        set_font(cells[1].paragraphs[0].add_run(result), size=9.5)
    borders(table)

    doc.add_page_break()
    add_heading(doc, "Prolog Reasoning")
    add_body(doc, "Prolog is MARLIN's rule-based logical thinker. It is used when an answer should follow explicit facts and rules instead of being guessed by the language model.")
    add_heading(doc, "Example Facts", 2)
    add_code(doc, [
        "task(task_report).",
        "urgent(task_report).",
        "belongs_to(task_report, project_marlin).",
        "active_project(project_marlin).",
        "not_completed(task_report).",
    ])
    add_heading(doc, "Example Rule", 2)
    add_code(doc, [
        "high_priority(Task) :-",
        "    task(Task),",
        "    urgent(Task),",
        "    belongs_to(Task, Project),",
        "    active_project(Project),",
        "    not_completed(Task).",
    ])
    add_body(doc, "The result is explainable: task_report is high priority because it is urgent, unfinished, and belongs to an active project. Python is the only component allowed to add facts or run predefined Prolog queries.")

    add_heading(doc, "Brain Graph")
    add_body(doc, "The graph visualises indexed content under C:\\Projects\\Projects and makes project structure easier to understand than a normal folder list.")
    add_bullets(doc, [
        "Large coloured nodes are main projects.",
        "Medium nodes are folders and small nodes are files.",
        "Solid links mean that a file or folder is contained in another folder.",
        "Cross-links connect folders that share file types.",
        "Dragging a node highlights its direct relationships.",
    ])

    doc.add_page_break()
    add_heading(doc, "MARLIN Capabilities")
    add_bullets(doc, [
        "Typed and English voice conversation with the Hey MARLIN wake phrase.",
        "Open accessible files, folders, websites, installed applications, and the Windows Camera.",
        "Play local videos, YouTube media, and control Windows media keys.",
        "Search indexed file names and short text snippets.",
        "Create files and folders and preview changes that could cause data loss.",
        "Set, display, complete, and snooze local alarms and reminders.",
        "Remember recent files, applications, projects, actions, and conversation context.",
        "Display and explain the brain graph and use Prolog for task recommendations.",
    ])
    add_heading(doc, "Privacy and Safety")
    add_body(doc, "The backend binds only to 127.0.0.1. Qwen, Whisper, Vosk, Piper, SQLite, and Prolog run locally. Opening and reading are direct actions. Closing applications and changing, moving, replacing, or deleting files require confirmation because they can cause data loss. Arbitrary shell commands remain blocked.")

    add_heading(doc, "Six Member Presentation Split")
    members = [
        "Member 1 - Problem and objective: introduce MARLIN and explain why a local personal assistant is useful.",
        "Member 2 - Architecture: explain the Python command router and the three service routes.",
        "Member 3 - Local AI and voice: explain Qwen, Ollama, Faster-Whisper, Vosk, and Piper.",
        "Member 4 - Database and graph: explain SQLite, indexing, projects, files, and relationships.",
        "Member 5 - Prolog: demonstrate facts, rules, high-priority tasks, and explanations.",
        "Member 6 - Actions and demonstration: show camera, applications, videos, reminders, safety, and shutdown.",
    ]
    add_numbered(doc, members)

    add_heading(doc, "Suggested Live Demonstration")
    add_numbered(doc, [
        "Run py main.py and show the desktop cockpit.",
        "Say or type open camera.",
        "Type explain my graph and show the project relationships.",
        "Type high priority tasks and point out the Prolog activity.",
        "Open a file or application with a natural command.",
        "Set a reminder and show its local storage details.",
        "Finish with turn yourself off to demonstrate lifecycle control.",
    ])
    add_heading(doc, "Short Presentation Summary")
    add_body(doc, "MARLIN is a fully local Windows AI assistant. Qwen handles conversation, Python controls actions, SQLite stores memory, Prolog performs logical reasoning, and the brain graph visualises project knowledge. This design allows MARLIN to understand, reason, remember, explain, and perform real computer tasks while protecting risky operations.")

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
